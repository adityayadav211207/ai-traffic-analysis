import os
import pandas as pd
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

from analytics.kpi import calculate_kpis
from analytics.smart_score import rank_high_risk_corridors


def _build_recommendation_rows(df):
    if df is None or df.empty:
        return [
            {"Priority": "High", "Action": "No operational dataset available for recommendations."},
        ]

    col_map = {str(c).lower().strip(): c for c in df.columns}
    road_col = col_map.get("road") or col_map.get("road_name") or col_map.get("location")
    city_col = col_map.get("city")
    severity_col = col_map.get("severity") or col_map.get("accident_severity")
    weather_col = col_map.get("weather")

    summary = []
    if road_col:
        top_roads = df[road_col].dropna().astype(str).value_counts().head(3)
        for road, count in top_roads.items():
            summary.append({
                "Priority": "High",
                "Action": f"Deploy additional traffic monitoring and enforcement on {road} to reduce recurring incidents ({count} recorded events).",
                "Owner": "Road Safety Team",
            })

    if severity_col:
        fatal_count = df[severity_col].astype(str).str.lower().str.contains("fatal|death|killed", case=False, na=False).sum()
        if fatal_count:
            summary.append({
                "Priority": "Critical",
                "Action": "Prioritize fatality hotspot mitigation, including signal timing optimization and emergency response review.",
                "Owner": "Operations Leadership",
            })

    if weather_col:
        weather_count = df[weather_col].dropna().astype(str).str.lower().value_counts().idxmax()
        if weather_count:
            summary.append({
                "Priority": "Medium",
                "Action": f"Review weather-responsive traffic controls and signage for {weather_count.title()} conditions.",
                "Owner": "Planning & Safety",
            })

    if not summary:
        summary.append({
            "Priority": "Medium",
            "Action": "Complete quarterly review of accident patterns and validate update cadence for operational datasets.",
            "Owner": "Management Office",
        })

    return summary[:6]


def generate_excel_report(df, output_path=None):
    """
    Generates a board-ready executive Excel workbook with clear KPI summaries,
    corridor rankings, operational recommendations, and telemetry data.
    """
    if output_path is None:
        os.makedirs("reports/excel", exist_ok=True)
        filename = f"Traffic_Executive_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        output_path = os.path.join("reports", "excel", filename)
    else:
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    kpis = calculate_kpis(df) if df is not None else {
        "total_accidents": 0, "fatal_accidents": 0, "high_risk_roads": 0,
        "total_cities": 0, "vehicle_types": 0, "total_injuries": 0, "traffic_score": 0
    }

    corridors = rank_high_risk_corridors(df, top_n=15) if df is not None else []
    recommendations = _build_recommendation_rows(df)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary_df = pd.DataFrame([
            {"Metric": "Total Recorded Incidents", "Value": kpis.get("total_accidents", 0), "Benchmark": "Operational baseline"},
            {"Metric": "Fatal Collisions", "Value": kpis.get("fatal_accidents", 0), "Benchmark": "Safety critical"},
            {"Metric": "Total Casualties / Injuries", "Value": kpis.get("total_injuries", 0), "Benchmark": "Impact severity"},
            {"Metric": "High-Risk Corridors", "Value": kpis.get("high_risk_roads", 0), "Benchmark": "Risk exposure"},
            {"Metric": "Cities Covered", "Value": kpis.get("total_cities", 0), "Benchmark": "Operational footprint"},
            {"Metric": "Vehicle Classes", "Value": kpis.get("vehicle_types", 0), "Benchmark": "Fleet diversity"},
            {"Metric": "Safety Performance Index", "Value": f"{kpis.get('traffic_score', 0)}%", "Benchmark": "Target > 80%"},
            {"Metric": "Report Generated On", "Value": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "Benchmark": "Executive review"},
        ])
        summary_df.to_excel(writer, sheet_name="Executive Summary", index=False)

        if corridors:
            corridor_df = pd.DataFrame(corridors)
            corridor_df = corridor_df.rename(columns={
                "road": "Road",
                "city": "City",
                "incidents": "Incidents",
                "fatalities": "Fatalities",
                "injuries": "Injuries",
                "safety_score": "Safety Score",
                "danger_rating": "Risk Level",
            })
            corridor_df.to_excel(writer, sheet_name="Risk Corridor Rankings", index=False)
        else:
            pd.DataFrame([{"Notice": "No corridor risk data available for this dataset."}]).to_excel(writer, sheet_name="Risk Corridor Rankings", index=False)

        rec_df = pd.DataFrame(recommendations)
        rec_df = rec_df.rename(columns={
            "Priority": "Priority",
            "Action": "Recommended Action",
            "Owner": "Owner",
        })
        rec_df.to_excel(writer, sheet_name="Operational Recommendations", index=False)

        if df is not None and not df.empty:
            df.head(500).to_excel(writer, sheet_name="Incident Telemetry Logs", index=False)

        workbook = writer.book
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        accent_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
        warning_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        critical_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        cell_font = Font(name="Segoe UI", size=10)
        title_font = Font(name="Segoe UI", size=12, bold=True, color="0F172A")
        thin_border = Border(
            left=Side(style="thin", color="D1D5DB"),
            right=Side(style="thin", color="D1D5DB"),
            top=Side(style="thin", color="D1D5DB"),
            bottom=Side(style="thin", color="D1D5DB")
        )

        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            sheet.sheet_view.showGridLines = True

            for row in sheet.iter_rows():
                for cell in row:
                    cell.border = thin_border
                    cell.font = cell_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            if sheet.title == "Executive Summary":
                sheet["A1"] = "Report Metric"
                sheet["B1"] = "Value"
                sheet["C1"] = "Context"
                for cell in sheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font

                sheet["A10"] = "Leadership View"
                sheet["A10"].font = title_font
                sheet["A10"].fill = accent_fill
                sheet.merge_cells("A10:C10")
                sheet["A11"] = "Use the risk ranking and recommendations tabs to prioritize operational action."
                sheet.merge_cells("A11:C11")
                sheet["A11"].alignment = Alignment(wrap_text=True, vertical="top")

                chart = BarChart()
                chart.type = "bar"
                chart.style = 10
                chart.title = "Executive KPI Snapshot"
                chart.y_axis.title = "Metric"
                chart.x_axis.title = "Value"
                values = Reference(sheet, min_col=2, min_row=1, max_row=7)
                labels = Reference(sheet, min_col=1, min_row=2, max_row=7)
                chart.add_data(values, titles_from_data=True)
                chart.set_categories(labels)
                chart.height = 7
                chart.width = 13
                sheet.add_chart(chart, "E2")

            if sheet.title == "Risk Corridor Rankings":
                for cell in sheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                for row in range(2, sheet.max_row + 1):
                    risk_cell = sheet.cell(row=row, column=7)
                    risk = str(risk_cell.value or "").lower()
                    if "critical" in risk:
                        risk_cell.fill = critical_fill
                    elif "high" in risk:
                        risk_cell.fill = warning_fill

            if sheet.title == "Operational Recommendations":
                for row in range(2, sheet.max_row + 1):
                    priority = str(sheet.cell(row=row, column=1).value or "").lower()
                    if priority == "critical":
                        sheet.cell(row=row, column=1).fill = critical_fill
                    elif priority == "high":
                        sheet.cell(row=row, column=1).fill = warning_fill

            if sheet.title == "Operational Recommendations":
                for cell in sheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font

            if sheet.title == "Incident Telemetry Logs":
                for cell in sheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font

            for col_idx, col in enumerate(sheet.columns, 1):
                max_len = 0
                for cell in col:
                    value = str(cell.value) if cell.value is not None else ""
                    if len(value) > max_len:
                        max_len = len(value)
                width = max(max_len + 3, 16)
                sheet.column_dimensions[get_column_letter(col_idx)].width = min(width, 30)

    return output_path
